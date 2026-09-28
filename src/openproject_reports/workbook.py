"""Excel workbook with stable tabular schema and restrained formatting."""

import io
from datetime import date, datetime
from importlib.metadata import version
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from openproject_reports.reporting import FIELDS, progress, projects

PROJECT_FIELDS = (
    "project_id",
    "project",
    "total_work_packages",
    "open",
    "in_progress",
    "blocked",
    "completed",
    "completed_last_7_days",
    "completed_last_30_days",
)
PROGRESS_FIELDS = (
    "completed_week",
    "completed",
    "cumulative_completions",
    "remaining_open_as_of_run",
    "snapshot_date",
)


def _table(
    book: Workbook, title: str, fields: tuple[str, ...], rows: list[dict[str, Any]]
) -> None:
    sheet = book.create_sheet(title)
    sheet.append(fields)
    for row in rows:
        values = []
        for field in fields:
            value = row.get(field)
            if isinstance(value, datetime):
                value = value.replace(tzinfo=None)
            values.append(value)
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, field in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(field) + 3, 15), 42
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"


def build_workbook(
    rows: list[dict[str, Any]],
    generated_at: datetime,
    base_url: str,
    timezone: str,
    drive_file_id: str | None = None,
) -> bytes:
    book = Workbook()
    book.remove(book.active)
    today = generated_at.date()
    project_rows = projects(rows, today)
    _table(book, "Work Packages", FIELDS, rows)
    _table(book, "Projects", PROJECT_FIELDS, project_rows)
    _table(book, "Progress", PROGRESS_FIELDS, progress(rows, today))
    try:
        repository_version = version("openproject-reports")
    except Exception:
        repository_version = "0.1.0"
    metadata = [
        ("generated_timestamp", generated_at.isoformat()),
        ("source_system", "OpenProject"),
        ("source_endpoint", f"{base_url.rstrip('/')}/api/v3/work_packages"),
        ("repository_version", repository_version),
        ("report_schema_version", "1"),
        ("timezone", timezone),
        ("row_count", len(rows)),
        ("project_count", len(project_rows)),
        ("google_drive_file_id", drive_file_id or ""),
    ]
    _table(
        book,
        "Metadata",
        ("key", "value"),
        [dict(zip(("key", "value"), pair, strict=True)) for pair in metadata],
    )
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
