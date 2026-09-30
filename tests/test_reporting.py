# mypy: disable-error-code="no-untyped-def,no-untyped-call,import-untyped,type-arg"
import json
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from openproject_reports.config import load_config
from openproject_reports.dagster.assets import openproject_reporting_dataset
from openproject_reports.dagster.definitions import defs
from openproject_reports.dagster.resources import RuntimeResource
from openproject_reports.openproject import OpenProjectClient
from openproject_reports.reporting import dataset, normalize, progress, projects
from openproject_reports.workbook import build_workbook

NOW = datetime(2026, 9, 28, 12, tzinfo=UTC)


def package(
    identifier: int, status: str, updated: str = "2026-09-20T00:00:00Z"
) -> dict:
    return {
        "id": identifier,
        "subject": f"Work {identifier}",
        "createdAt": "2026-09-01T00:00:00Z",
        "updatedAt": updated,
        "closedAt": "2026-09-22T00:00:00Z" if status == "Closed" else None,
        "_links": {
            "project": {"href": "/api/v3/projects/2", "title": "Project B"},
            "status": {"title": status},
            "type": {"title": "Story"},
        },
    }


class Response:
    def __init__(self, data):
        self.data = data

    def raise_for_status(self):
        pass

    def json(self):
        return self.data


class Session:
    def __init__(self):
        self.offsets = []

    def get(self, url, params, timeout):
        assert url.endswith("/work_packages")
        assert timeout == 60
        assert params["filters"] == "[]"
        assert params["sortBy"] == '[["id","asc"]]'
        self.offsets.append(params["offset"])
        identifier = params["offset"]
        return Response(
            {"total": 2, "_embedded": {"elements": [package(identifier, "Closed")]}}
        )


def test_pagination_retains_completed():
    session = Session()
    rows = OpenProjectClient("https://example.test", "token", session).work_packages(1)  # type: ignore[arg-type]
    assert session.offsets == [1, 2]
    assert [row["id"] for row in rows] == [1, 2]


def test_project_catalog_paginates_and_excludes_non_projects():
    class ProjectSession(Session):
        def get(self, url, params, timeout):
            assert url.endswith("/projects")
            assert timeout == 60
            assert params["filters"] == "[]"
            self.offsets.append(params["offset"])
            identifier = params["offset"]
            item_type = "Project" if identifier == 1 else "Program"
            return Response(
                {
                    "total": 2,
                    "_embedded": {
                        "elements": [
                            {
                                "id": identifier,
                                "name": f"P{identifier}",
                                "_type": item_type,
                            }
                        ]
                    },
                }
            )

    session = ProjectSession()
    rows = OpenProjectClient("https://example.test", "token", session).projects(1)  # type: ignore[arg-type]
    assert session.offsets == [1, 2]
    assert [row["name"] for row in rows] == ["P1"]


def test_dataset_rollups_and_workbook():
    rows = dataset(
        [package(2, "Closed"), package(1, "In Progress"), package(2, "Closed")],
        "https://example.test",
        NOW,
        "America/New_York",
    )
    assert [row["work_package_id"] for row in rows] == [1, 2]
    assert rows[1]["is_complete"] is True
    assert rows[1]["days_to_complete"] == 21
    assert rows[1]["completed_week"] == "2026-W39"
    assert rows[0]["is_open"] is True
    assert projects(rows, NOW.date())[0]["completed"] == 1
    assert (
        projects(
            rows,
            NOW.date(),
            [{"id": 2, "name": "Project B"}, {"id": 3, "name": "Empty"}],
        )[1]["total_work_packages"]
        == 0
    )
    assert progress(rows, NOW.date())[0]["cumulative_completions"] == 1
    book = load_workbook(
        BytesIO(build_workbook(rows, NOW, "https://example.test", "America/New_York"))
    )
    assert book.sheetnames == ["Work Packages", "Projects", "Progress", "Metadata"]
    assert book["Work Packages"].max_row == 3
    catalog_book = load_workbook(
        BytesIO(
            build_workbook(
                rows,
                NOW,
                "https://example.test",
                "America/New_York",
                project_catalog=[
                    {"id": 2, "name": "Project B"},
                    {"id": 3, "name": "Empty"},
                ],
            )
        )
    )
    assert catalog_book["Projects"].max_row == 3
    assert book["Work Packages"].freeze_panes == "A2"
    assert book["Work Packages"].auto_filter.ref == book["Work Packages"].dimensions


@pytest.mark.parametrize(
    ("status", "is_complete", "is_open"),
    [
        ("Closed", True, False),
        ("Done", True, False),
        ("In Progress", False, True),
        ("Blocked", False, True),
        ("rejected", False, False),
        ("Rejected", False, False),
        ("REJECTED", False, False),
    ],
)
def test_status_classification(status, is_complete, is_open):
    row = dataset([package(1, status)], "https://example.test", NOW, "UTC")[0]
    assert row["status"] == status
    assert row["is_complete"] is is_complete
    assert row["is_open"] is is_open


def test_rejected_remains_in_work_packages_but_not_rollups():
    rows = dataset(
        [package(1, "In Progress"), package(2, "Closed"), package(3, "Rejected")],
        "https://example.test",
        NOW,
        "UTC",
    )
    assert [row["status"] for row in rows] == ["In Progress", "Closed", "Rejected"]
    project = projects(rows, NOW.date())[0]
    assert project["total_work_packages"] == 3
    assert project["open"] == 1
    assert project["completed"] == 1
    assert progress(rows, NOW.date())[0]["remaining_open_as_of_run"] == 1

    book = load_workbook(
        BytesIO(build_workbook(rows, NOW, "https://example.test", "UTC"))
    )
    work_packages = book["Work Packages"]
    headers = [cell.value for cell in work_packages[1]]
    rejected = dict(
        zip(headers, (cell.value for cell in work_packages[4]), strict=True)
    )
    assert work_packages.max_row == 4
    assert rejected["status"] == "Rejected"
    assert rejected["is_open"] is False
    assert rejected["is_complete"] is False


def test_config_validation():
    with pytest.raises(ValueError, match="Missing runtime settings"):
        load_config({})
    settings = load_config(
        {
            "OPENPROJECT_BASE_URL": "https://example.test/",
            "OPENPROJECT_API_TOKEN": "secret",
            "GOOGLE_DRIVE_CREDENTIALS_JSON": json.dumps(
                {
                    "type": "authorized_user",
                    "client_id": "a",
                    "client_secret": "b",
                    "refresh_token": "c",
                }
            ),
            "GOOGLE_DRIVE_FOLDER_ID": "folder",
        }
    )
    assert settings.openproject_base_url == "https://example.test"
    with pytest.raises(ValueError, match="My Drive needs user OAuth"):
        load_config(
            {
                "OPENPROJECT_BASE_URL": "https://example.test",
                "OPENPROJECT_API_TOKEN": "secret",
                "GOOGLE_DRIVE_CREDENTIALS_JSON": json.dumps(
                    {"type": "service_account"}
                ),
                "GOOGLE_DRIVE_FOLDER_ID": "folder",
            }
        )


def test_definitions_registered():
    assert len(defs.resolve_asset_graph().get_all_asset_keys()) == 5
    assert defs.get_job_def("openproject_full_refresh")
    assert defs.get_schedule_def("openproject_daily")


def typed_package(identifier, kind, status, parent=None):
    item = package(identifier, status)
    item["_links"]["type"]["title"] = kind
    if parent is not None:
        item["_links"]["parent"] = {"href": f"/api/v3/work_packages/{parent}"}
    return item


def test_story_hierarchy_and_operational_counts():
    items = [
        typed_package(10, "Initiative", "Closed"),
        typed_package(11, "Epic", "In Progress", 10),
        typed_package(12, "Story", "In Progress", 11),
        typed_package(13, "Story", "Closed", 11),
        typed_package(14, "Story", "Blocked", 11),
        typed_package(15, "Task", "Closed", 12),
        typed_package(16, "Milestone", "Blocked"),
    ]
    items[2]["_links"]["version"] = {"href": "/api/v3/versions/20", "title": "R1"}
    rows = dataset(items, "https://example.test", NOW, "UTC")
    assert [row["work_package_id"] for row in rows] == [12, 13, 14]
    assert rows[0]["initiative_id"] == 10
    assert rows[0]["initiative"] == "Work 10"
    assert rows[0]["epic_id"] == 11
    assert rows[0]["epic"] == "Work 11"
    assert rows[0]["release_id"] == 20
    assert rows[0]["release"] == rows[0]["version"] == "R1"
    assert rows[0]["url"] == "https://example.test/work_packages/12"
    project = projects(rows, NOW.date())[0]
    assert (
        project["total_work_packages"],
        project["open"],
        project["in_progress"],
        project["blocked"],
        project["completed"],
    ) == (3, 2, 1, 1, 1)
    assert progress(rows, NOW.date())[0]["completed"] == 1
    for item in items[:2]:
        item["_links"]["status"]["title"] = "Blocked"
    assert projects(dataset(items, "https://example.test", NOW, "UTC"), NOW.date()) == [
        project
    ]


def test_missing_and_cyclic_hierarchy():
    orphan = typed_package(1, "Story", "New", 99)
    cycle = typed_package(2, "Epic", "New", 3)
    story = typed_package(3, "Story", "New", 2)
    rows = dataset([orphan, cycle, story], "https://example.test", NOW, "UTC")
    assert rows[0]["initiative_id"] is None
    assert rows[0]["epic_id"] is None
    assert rows[1]["epic_id"] == 2
    assert rows[1]["initiative_id"] is None


def activity(identifier, comment, created="2026-09-22T00:00:00Z"):
    return {"id": identifier, "createdAt": created, "comment": {"raw": comment}}


def test_latest_implementation_summary_and_workbook():
    activities = [
        activity(5, "Routine discussion", "2026-09-25T00:00:00Z"),
        activity(3, "# Implementation summary\nFinal changes"),
        activity(1, "Implementation update (WP-1)\nEarlier", "2026-09-20T00:00:00Z"),
        activity(2, "Implementation summary\nOlder tie"),
        activity(6, ""),
        {"id": 7, "comment": "invalid"},
        {"id": 8, "comment": {"raw": None}},
    ]
    rows = dataset(
        [package(1, "Closed"), package(2, "New")],
        "https://example.test",
        NOW,
        "UTC",
        activities={1: activities},
    )
    assert (
        rows[0]["implementation_summary"] == "# Implementation summary\nFinal changes"
    )
    assert rows[1]["implementation_summary"] is None
    book = load_workbook(
        BytesIO(build_workbook(rows, NOW, "https://example.test", "UTC"))
    )
    headers = [cell.value for cell in book["Work Packages"][1]]
    assert (
        book["Work Packages"].cell(2, headers.index("implementation_summary") + 1).value
        == rows[0]["implementation_summary"]
    )
    assert "Routine discussion" not in str(list(book["Work Packages"].values))


def test_activity_pagination_and_empty_collection():
    class ActivitySession(Session):
        def get(self, url, params, timeout):
            assert url.endswith("/work_packages/1/activities")
            assert timeout == 60
            self.offsets.append(params["offset"])
            return Response(
                {
                    "total": 2,
                    "_embedded": {
                        "elements": [
                            activity(params["offset"], "Implementation update\nChange")
                        ]
                    },
                }
            )

    session = ActivitySession()
    client = OpenProjectClient("https://example.test", "token", session)  # type: ignore[arg-type]
    result = client.activities(1, 1)
    assert [row["id"] for row in result] == [1, 2]
    assert session.offsets == [1, 2]

    class EmptySession(Session):
        def get(self, url, params, timeout):
            assert url.endswith("/activities")
            assert params["pageSize"] == 100
            assert timeout == 60
            return Response({"total": 0, "_embedded": {"elements": []}})

    client = OpenProjectClient("https://example.test", "token", EmptySession())  # type: ignore[arg-type]
    assert client.activities(1) == []


def test_operational_consumers_exclude_hierarchy_rows():
    rows = [
        normalize(typed_package(i, kind, "Closed"), "https://example.test", NOW, "UTC")
        for i, kind in enumerate(["Story", "Epic", "Initiative"], 1)
    ]
    assert projects(rows, NOW.date())[0]["completed"] == 1
    assert progress(rows, NOW.date())[0]["completed"] == 1
    book = load_workbook(
        BytesIO(build_workbook(rows, NOW, "https://example.test", "UTC"))
    )
    assert book["Work Packages"].max_row == 2


def test_reporting_asset_fetches_only_story_activities(monkeypatch):
    requested = []

    class Client:
        def activities(self, identifier):
            requested.append(identifier)
            return [activity(1, "Implementation update (WP-3)\nImplemented")]

    class Settings:
        openproject_base_url = "https://example.test"
        timezone = "UTC"

    monkeypatch.setattr(RuntimeResource, "settings", lambda _: Settings())
    monkeypatch.setattr(RuntimeResource, "openproject", lambda _: Client())
    rows = openproject_reporting_dataset(
        runtime=RuntimeResource(),
        openproject_work_packages=[
            typed_package(1, "Initiative", "Closed"),
            typed_package(2, "Epic", "Blocked", 1),
            typed_package(3, "Story", "Closed", 2),
        ],
    )
    assert requested == [3]
    assert isinstance(rows, list)
    assert len(rows) == 1
    assert (
        rows[0]["implementation_summary"] == "Implementation update (WP-3)\nImplemented"
    )
    assert rows[0]["initiative_id"] == 1
    assert rows[0]["epic_id"] == 2
